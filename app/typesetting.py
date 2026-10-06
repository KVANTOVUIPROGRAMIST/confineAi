"""Trusted Markdown-to-LaTeX conversion and bounded, untrusted compilation."""
import hashlib
import io
import json
import logging
import os
from pathlib import Path
import platform
import re
import subprocess
import tarfile
import tempfile
import threading
import zipfile

import httpx
from fastapi import HTTPException
from pypdf import PdfReader, PdfWriter

from .config import ROOT

VERSION = '0.17.0'
ASSETS = {
    ('Windows', 'AMD64'): ('x86_64-pc-windows-msvc.zip', 'f61ce51f0b0ade1015b7de7ef368541c5424e9756ecbd0d7af97d6d48030845f'),
    ('Linux', 'x86_64'): ('x86_64-unknown-linux-musl.tar.gz', '8533d07f9ccbd7a65824b9e0459041bca34af1eb33daba48f59215593753a3b7'),
}
LOCK = threading.Lock()
COMPILE_LOCK = threading.Lock()
READER = 'markdown-raw_tex-raw_html-raw_attribute+tex_math_dollars'
MATH_COMMANDS = set(r'''frac dfrac tfrac cfrac binom dbinom tbinom sqrt sum prod coprod int iint iiint oint lim min max inf sup log ln exp sin cos tan sec csc cot arcsin arccos arctan sinh cosh tanh gcd det Pr mod bmod pmod left right middle big Big bigg Bigg bigl bigr Bigl Bigr biggl biggr Biggl Biggr overline underline underbrace overbrace overset underset stackrel text textrm textnormal textbf textit mathrm mathbf mathit mathsf mathtt mathcal mathbb mathscr boldsymbol operatorname boxed phantom vphantom hphantom substack limits nolimits displaystyle textstyle scriptstyle scriptscriptstyle quad qquad hspace vspace tag notag nonumber label ref eqref begin end alpha beta gamma delta epsilon varepsilon zeta eta theta vartheta iota kappa lambda mu nu xi pi varpi rho varrho sigma varsigma tau upsilon phi varphi chi psi omega Gamma Delta Theta Lambda Xi Pi Sigma Upsilon Phi Psi Omega forall exists nexists in notin ni subset supset subseteq supseteq nsubseteq nsupseteq cup cap bigcup bigcap setminus emptyset varnothing land lor lnot neg wedge vee implies iff to mapsto rightarrow leftarrow leftrightarrow Rightarrow Leftarrow Leftrightarrow longrightarrow longleftarrow longleftrightarrow Longrightarrow Longleftarrow Longleftrightarrow uparrow downarrow le leq ge geq ne neq approx sim simeq equiv cong propto pm mp times div cdot cdots ldots dots vdots ddots ast star circ bullet infinity infty partial nabla ell hbar Re Im angle triangle parallel perp mid vert Vert lvert rvert lVert rVert langle rangle lceil rceil lfloor rfloor lbrace rbrace backslash colon choose atop over brace brack dot ddot hat widehat tilde widetilde vec bar breved check acute grave space newline nobreak thinspace medspace thickspace mathop mathrel mathbin mathord mathopen mathclose mathpunct cancel not top bot percent degree'''.split())
MATH_ENVIRONMENTS = {'aligned', 'alignedat', 'gathered', 'split', 'cases', 'matrix', 'pmatrix', 'bmatrix', 'Bmatrix', 'vmatrix', 'Vmatrix', 'array', 'smallmatrix'}


class TypesetError(HTTPException):
    def __init__(self, detail='The LaTeX document could not be compiled. Retry the saved file; your allowance was restored.'):
        super().__init__(502, detail)


def compiler():
    """Install a pinned, checksum-verified application compiler, never system TeX."""
    target = ROOT / '.tools' / ('tectonic.exe' if os.name == 'nt' else 'tectonic')
    with LOCK:
        if target.exists():
            return str(target)
        asset, digest = ASSETS.get((platform.system(), platform.machine()), (None, None))
        if not asset:
            raise TypesetError('LaTeX compilation is not configured for this server platform.')
        url = f'https://github.com/tectonic-typesetting/tectonic/releases/download/tectonic%40{VERSION}/tectonic-{VERSION}-{asset}'
        response = httpx.get(url, follow_redirects=True, timeout=120)
        response.raise_for_status()
        data = response.content
        if hashlib.sha256(data).hexdigest() != digest:
            raise TypesetError('The LaTeX compiler download failed its integrity check.')
        # Extract only the named executable, never archive-controlled paths.
        if asset.endswith('.zip'):
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                candidates = [n for n in archive.namelist() if Path(n).name == target.name]
                executable = archive.read(candidates[0]) if len(candidates) == 1 else None
        else:
            with tarfile.open(fileobj=io.BytesIO(data), mode='r:gz') as archive:
                candidates = [n for n in archive.getmembers() if n.isfile() and Path(n.name).name == target.name]
                executable = archive.extractfile(candidates[0]).read() if len(candidates) == 1 else None
        if not executable:
            raise TypesetError('The LaTeX compiler archive did not contain the expected executable.')
        target.parent.mkdir(exist_ok=True)
        target.write_bytes(executable)
        target.chmod(0o755)
    return str(target)


def pandoc(text, to, extra_args=None, outputfile=None):
    import pypandoc
    result = pypandoc.convert_text(text, to, format=READER, extra_args=extra_args or [], outputfile=outputfile)
    return result.replace('\r\n', '\n')


def normalize_math(text):
    """Repair a mismatched single-dollar display terminator; never alter math content."""
    text = text.replace('\r\n', '\n')
    protected = [(m.start(), m.end()) for m in re.finditer(r'(?ms)^\s*(`{3,}|~{3,})[^\n]*\n.*?^\s*\1\s*$|`[^`\n]*`', text)]
    changes, state = [], None
    for match in re.finditer(r'\${1,2}', text):
        if any(start <= match.start() < end for start, end in protected):
            continue
        slashes = 0
        cursor = match.start() - 1
        while cursor >= 0 and text[cursor] == '\\':
            slashes += 1
            cursor -= 1
        if slashes % 2:
            continue
        token = match.group()
        if state is None:
            state = token
        elif state == '$$' and token == '$':
            end = match.end() + (1 if text[match.end():match.end()+1] == '|' else 0)
            changes.append((match.start(), end, '$$'))
            state = None
        elif token == state:
            state = None
        else:
            raise TypesetError('The generated mathematics has mismatched delimiters.')
    if state:
        raise TypesetError('The generated mathematics has an unclosed formula.')
    for start, end, replacement in reversed(changes):
        text = text[:start] + replacement + text[end:]
    return text


def validate_markdown(text):
    if len(text) > 300_000:
        raise TypesetError('The generated document exceeds the formatting limit.')
    text = normalize_math(text)
    tree = json.loads(pandoc(text, 'json'))

    def walk(node):
        if isinstance(node, dict):
            kind = node.get('t')
            if kind in ('Image', 'RawBlock', 'RawInline'):
                raise TypesetError('The generated document contains unsupported embedded content.')
            if kind == 'Math':
                math = node['c'][1]
                if '^^' in math or '\x00' in math or re.search(r'(?<!\\)%', math):
                    raise TypesetError('The generated mathematics contains unsafe formatting.')
                commands = re.findall(r'\\([a-zA-Z]+)', math)
                if set(commands) - MATH_COMMANDS:
                    raise TypesetError('The generated mathematics uses an unsupported LaTeX command.')
                for env in re.findall(r'\\(?:begin|end)\s*\{([^}]+)\}', math):
                    if env not in MATH_ENVIRONMENTS:
                        raise TypesetError('The generated mathematics uses an unsupported LaTeX environment.')
            for value in node.values():
                walk(value)
        elif isinstance(node, list):
            for value in node:
                walk(value)
    walk(tree)


def markdown_document(title, course, results, student_name='', layout='rebuild'):
    def literal(value):
        return re.sub(r'([\\`*_{}\[\]<>#$!|])', r'\\\1', str(value).replace('\n', ' '))
    lines = [f'# {literal(title)}', literal(f'{course.code} - {course.title}')]
    if student_name:
        lines.append(literal(student_name))
    for result in results:
        lines.append('## ' + literal(result['label']))
        if layout == 'rebuild':
            lines.append(normalize_math(result['question']))
            if result.get('uses_diagram'):
                lines.append('*The original figure is reproduced in the assignment appendix.*')
            lines.append('**Solution**')
        lines.append(normalize_math(result['answer']['final_answer']))
    return '\n\n'.join(lines) + '\n'


PREAMBLE = r'''\documentclass[11pt,letterpaper]{article}
\usepackage[margin=0.8in]{geometry}
\usepackage{amsmath,amssymb,mathtools}
\usepackage{mathrsfs,cancel}
\usepackage{longtable,booktabs,array}
\usepackage{graphicx,pdfpages}
\usepackage{fancyvrb}
\usepackage[unicode,hidelinks]{hyperref}
\setlength{\parindent}{0pt}
\setlength{\parskip}{6pt}
\setlength{\emergencystretch}{3em}
\setcounter{secnumdepth}{0}
\providecommand{\tightlist}{\setlength{\itemsep}{0pt}\setlength{\parskip}{0pt}}
\providecommand{\degree}{^{\circ}}
\providecommand{\percent}{\%}
\providecommand{\infinity}{\infty}
\providecommand{\breved}[1]{\breve{#1}}
\begin{document}
'''


def latex_document(markdown, include_original=False):
    markdown = normalize_math(markdown)
    validate_markdown(markdown)
    body = pandoc(markdown, 'latex', ['--syntax-highlighting=none', '--wrap=none'])
    if include_original:
        body += r'\clearpage\includepdf[pages=-,pagecommand={\thispagestyle{empty}}]{original.pdf}' + '\n'
    return PREAMBLE + body + '\n\\end{document}\n'


def compile_pdf(latex, original=None):
    executable = compiler()
    with COMPILE_LOCK, tempfile.TemporaryDirectory(prefix='confine-tex-') as directory:
        root = Path(directory)
        (root / 'submission.tex').write_text(latex, encoding='utf-8', newline='\n')
        if original:
            # Copy page content only: no annotations, embedded files or document actions.
            writer = PdfWriter()
            for page in PdfReader(io.BytesIO(original)).pages:
                copied = writer.add_page(page)
                copied.pop('/Annots', None)
            writer.write(root / 'original.pdf')
        try:
            run = subprocess.run([executable, '-X', 'compile', '--untrusted', '--reruns', '0', 'submission.tex'],
                cwd=root, capture_output=True, timeout=180,
                env={**os.environ, 'TECTONIC_UNTRUSTED_MODE': '1'}, check=False)
        except (OSError, subprocess.TimeoutExpired):
            raise TypesetError() from None
        pdf = root / 'submission.pdf'
        if run.returncode or not pdf.exists():
            # Compiler output may contain student content: never log or return it.
            raise TypesetError()
        data = pdf.read_bytes()
        reader = PdfReader(io.BytesIO(data))
        if not reader.pages or len(reader.pages) > 150 or len(data) > 20_000_000:
            raise TypesetError('The compiled document exceeds the output limit.')
        return data


def warm_typesetter():
    """Warm trusted runtime resources without delaying web-service health checks."""
    try:
        sample = r'''# Document

$$\sum_{k=1}^{n} k=\frac{n(n+1)}{2},\qquad \binom{n}{r},\quad \mathbb{R}$$

**Solution** with *emphasis* and an equation $x^2$.

```
print("ready")
```
'''
        compile_pdf(latex_document(sample))
    except Exception as exc:
        logging.getLogger(__name__).warning('Typesetter warmup failed (%s); next document can retry.', type(exc).__name__)


def render_completed(title, course, results, student_name, layout, original_file):
    markdown = markdown_document(title, course, results, student_name, layout)
    include_original = layout == 'rebuild' and original_file.mime == 'application/pdf' and any(r.get('uses_diagram') for r in results)
    latex = latex_document(markdown, include_original)
    pdf = compile_pdf(latex, original_file.data if include_original else None)
    with tempfile.TemporaryDirectory(prefix='confine-word-') as directory:
        path = Path(directory) / 'submission.docx'
        # Pandoc creates native Word equations (OMML), preserving editable mathematics.
        pandoc(markdown, 'docx', ['--standalone'], str(path))
        if include_original:
            append_word_figures(path, original_file.data)
        docx = path.read_bytes()
        if len(docx) > 20_000_000:
            raise TypesetError('The Word document exceeds the output limit. Use a smaller assignment file.')
    return pdf, markdown, latex, docx, include_original


def append_word_figures(path, original):
    import pypdfium2 as pdfium
    from docx import Document
    from docx.shared import Inches
    document = Document(path)
    with pdfium.PdfDocument(original) as source:
        for index in range(len(source)):
            document.add_page_break()
            if index == 0:
                document.add_heading('Original assignment figures and instructions', 1)
            page = source[index]
            try:
                width, height = page.get_size()
                scale = min(2, 1500 / max(width, height))
                bitmap = page.render(scale=scale, draw_annots=False, may_draw_forms=False)
                try:
                    image = bitmap.to_pil()
                    stream = io.BytesIO()
                    image.save(stream, format='PNG')
                    image.close()
                finally:
                    bitmap.close()
                stream.seek(0)
                inches = min(6, 7.5 * width / height)
                document.add_picture(stream, width=Inches(inches))
            finally:
                page.close()
    document.save(path)


def latex_project(latex, original=None):
    output = io.BytesIO()
    with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('submission.tex', latex)
        if original:
            writer = PdfWriter()
            for page in PdfReader(io.BytesIO(original)).pages:
                copied = writer.add_page(page)
                copied.pop('/Annots', None)
            data = io.BytesIO()
            writer.write(data)
            archive.writestr('original.pdf', data.getvalue())
        archive.writestr('README.txt', 'Compile submission.tex with XeLaTeX or Tectonic. Keep original.pdf beside it when included.\n')
    return output.getvalue()
