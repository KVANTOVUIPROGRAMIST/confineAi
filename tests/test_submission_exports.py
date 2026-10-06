import io
from types import SimpleNamespace

from docx import Document
from pypdf import PdfReader

from app.documents import render_submission, submission_docx, submission_markdown


COURSE = SimpleNamespace(code='CSE 29', title='Systems Programming')


def answer(content, status='answered', **kwargs):
    return {'status': status, 'summary': 'INTERNAL AUDIT SUMMARY', 'steps': [],
            'final_answer': content, 'sources': [], **kwargs}


def pdf_text(data):
    return '\n'.join(page.extract_text() for page in PdfReader(io.BytesIO(data)).pages)


def test_submission_keeps_original_labels_and_hides_audit():
    results = [
        {'label': '4(a)', 'question': 'PRIVATE QUESTION TEXT',
         'answer': answer('The result is **0.375**.\n\nP(X=2) = 3 * (0.5)^2 * (0.5).',
                          steps=[{'explanation': 'PRIVATE STEP', 'source_ids': ['S7']}],
                          sources=[{'id': 'S7', 'name': 'Lecture 2', 'page': 3, 'text': 'PRIVATE SOURCE EXCERPT'}])},
        {'label': '4(b)', 'question': 'Different question', 'answer': answer('E[X] = n * p = 1.5.')},
    ]
    markdown = submission_markdown('Assignment 1', COURSE, results, 'Test Student')
    pdf = pdf_text(render_submission('Assignment 1', COURSE, results, 'Test Student'))
    doc = Document(io.BytesIO(submission_docx('Assignment 1', COURSE, results, 'Test Student')))
    word = '\n'.join(paragraph.text for paragraph in doc.paragraphs)
    for output in (markdown, pdf, word):
        assert '4(a)' in output and '4(b)' in output
        assert 'Test Student' in output
        assert '3 * (0.5)^2 * (0.5)' in output
        assert 'PRIVATE' not in output and 'INTERNAL AUDIT' not in output and 'S7' not in output
        assert 'Confine' not in output and 'Question 1' not in output
        assert 'Incomplete assignment' not in output


def test_submission_code_and_table_are_editable_and_safe():
    code = 'int main(void) {\n\tif (2 < 3 && 4 > 1) {\n        return 0;\n    }\n}'
    content = f'### Implementation\n\n```c\n{code}\n```\n\n| Input | Output |\n|---|---|\n| A & B | <literal> |\n\n2 < 3 and 4 > 1.\n<script>alert("safe")</script>'
    results = [{'label': 'Problem 8 - Part C', 'question': '', 'answer': answer(content)}]
    markdown = submission_markdown('Code <assignment>', COURSE, results)
    word = Document(io.BytesIO(submission_docx('Code <assignment>', COURSE, results)))
    assert code in markdown
    assert code in '\n'.join(paragraph.text for paragraph in word.paragraphs)
    assert word.tables[0].cell(1, 0).text == 'A & B'
    assert word.tables[0].cell(1, 1).text == '<literal>'
    pdf = pdf_text(render_submission('Code <assignment>', COURSE, results))
    assert '2 < 3 && 4 > 1' in pdf
    assert 'alert("safe")' in pdf
    assert '<literal>' in pdf


def test_incomplete_export_never_includes_withheld_final_answer():
    results = [{'label': '7(b)', 'question': 'Unsupported task',
                'answer': answer('UNSUPPORTED SOLUTION', status='needs_materials',
                                 summary='Upload the lecture covering normal approximations.')}]
    markdown = submission_markdown('Assignment 3', COURSE, results)
    pdf = pdf_text(render_submission('Assignment 3', COURSE, results))
    word = '\n'.join(p.text for p in Document(io.BytesIO(submission_docx('Assignment 3', COURSE, results))).paragraphs)
    for output in (markdown, pdf, word):
        assert 'Incomplete assignment' in output
        assert '7(b)' in output
        assert 'Upload the lecture' in output
        assert 'UNSUPPORTED SOLUTION' not in output


def test_internal_citation_is_human_readable_and_long_code_can_span_pages():
    content = 'This follows from the lecture [S8].\n\n```python\n' + '\n'.join(
        '    values.append("' + 'A' * 200 + '")' for _ in range(85)) + '\n```'
    results = [{'label': 'Appendix A', 'question': '',
                'answer': answer(content, sources=[{'id': 'S8', 'name': 'Lecture notes', 'page': 5, 'text': 'HIDDEN EXCERPT'}])}]
    pdf = render_submission('Long code example', COURSE, results)
    reader = PdfReader(io.BytesIO(pdf))
    text = '\n'.join(page.extract_text() for page in reader.pages)
    assert len(reader.pages) > 1
    assert '(Lecture notes, p. 5)' in text and 'S8' not in text and 'HIDDEN EXCERPT' not in text
    markdown = submission_markdown('Long code example', COURSE, results)
    assert 'A' * 200 in markdown
